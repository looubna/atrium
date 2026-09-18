"""Resume files: type detection by content, bounded text extraction, and the profile schema.

Uploaded files are untrusted. The type is decided by the bytes (not the name or
the declared content type), PDF extraction is limited in pages and characters,
encrypted PDFs are refused, and only plain text leaves this module.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Literal

from pydantic import BaseModel, Field

ResumeKind = Literal["pdf", "docx", "txt", "md"]

MAX_PAGES = 30
MAX_TEXT_CHARS = 60_000
MIN_TEXT_CHARS = 200
CONTENT_TYPES: dict[str, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
}


class UnsupportedResume(ValueError):
    """The file isn't a readable PDF, DOCX or text resume."""


def detect_kind(filename: str, data: bytes) -> ResumeKind:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
        raise UnsupportedResume("unsupported archive")
    if b"\x00" in data[:4096]:
        raise UnsupportedResume("binary file")
    try:
        data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedResume("text must be UTF-8") from exc
    return "md" if extension == "md" else "txt"


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def extract_text(kind: ResumeKind, data: bytes) -> str:
    if kind == "pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise UnsupportedResume("encrypted PDF")
            parts: list[str] = []
            for page in reader.pages[:MAX_PAGES]:
                parts.append(page.extract_text() or "")
                if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                    break
        except PdfReadError as exc:
            raise UnsupportedResume("unreadable PDF") from exc
        text = "\n".join(parts)
    elif kind == "docx":
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # python-docx raises several unrelated types for corrupt files
            raise UnsupportedResume("unreadable DOCX") from exc
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(lines)
    else:
        text = data.decode("utf-8")

    text = _normalize(text)
    if len(text) < MIN_TEXT_CHARS:
        raise UnsupportedResume("no readable text (scanned PDFs are not supported)")
    return text


class RoleEntry(BaseModel):
    title: str
    organization: str
    period: str
    highlights: list[str]


class ResumeProfile(BaseModel):
    summary: str = Field(description="Two or three sentences describing the candidate's background")
    currentTitle: str
    yearsExperience: int = Field(description="Approximate total years of professional experience, 0 if unclear")
    roles: list[RoleEntry]
    skills: list[str]
    achievements: list[str] = Field(description="Concrete, ideally quantified accomplishments")
    education: list[str]


def fake_profile(text: str) -> ResumeProfile:
    """Deterministic stand-in for local development and tests."""
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
    return ResumeProfile(
        summary=text[:280].replace("\n", " "),
        currentTitle=first_line[:80],
        yearsExperience=0,
        roles=[],
        skills=sorted({w.strip(",.;:").lower() for w in text.split() if w.istitle() and len(w) > 3})[:12],
        achievements=[line.strip()[:200] for line in text.splitlines() if re.search(r"\d+%|\$\d", line)][:5],
        education=[],
    )
